"""
Места лова: предложение объединить разные написания одного места и применение подтверждённой таблицы.

    python tools/place_groups.py suggest --password … [--out places_groups.xlsx]
    python tools/place_groups.py apply   --password … --file places_groups.xlsx [--dry-run]

suggest — читает справочник «Места лова» и число строк уловов по каждому написанию, группирует похожие
написания внутри водоёма (одинаковые номера и вид места: научная станция, насосная станция, промысловый
участок, устье реки, залив, мыс; близкое написание названия) и сохраняет таблицу для проверки ихтиологом.
В столбце «Решение» ихтиолог оставляет «объединить» или пишет «отдельно», при необходимости правит
«Название места». Группы, которые нужно проверить особенно, помечены в «Примечании».

apply — для каждой группы с решением «объединить» (или каждой группы из services/data_admin/app/seed/place_groups.json —
написания, уже подтверждённые заказчиком): записи водоёма с написаниями группы сводятся в одну — название из таблицы,
остальные написания становятся «другими написаниями», лишние записи удаляются (история сохраняется).
Записи ищутся по водоёму и написанию, поэтому таблицу можно применить и на другой установке.
Загруженные данные не меняются: новые загрузки и правки приводят написания к названию места.
"""
from __future__ import annotations

import argparse
import difflib
import getpass
import json
import re
import sys
from collections import Counter, defaultdict

import httpx
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

KINDS = [  # (вид места, признаки в написании, каноническое название с номером)
    ("pump", r"насос", "насосная станция №{n}"),
    ("prom", r"пром", "промысловый участок №{n}"),
    ("sci", r"науч", "научная станция №{n}"),
    ("st", r"станц|\bст\b", "научная станция №{n}"),
    ("mouth", r"\bусть|\bуст\b", None),
    ("bay", r"залив|\bзал\b", None),
]
STOP = r"\b(м|мыс|р|река|уч|участок|п|пос|оз|озеро|озера|о|остров|ост|разрез|вдхр|станция|ст|науч|насосная|насосный|" \
       r"насосной|насосн|пром|устье|устья|уст|залив|зал|зона|отдыха)\b"


def norm(s: str) -> str:
    return re.sub(r"\s+", " ", str(s).lower().replace("ё", "е")).strip()


def features(name: str):
    s = norm(name)
    kind = next((k for k, rx, _ in KINDS if re.search(rx, s)), None)
    nums = tuple(re.findall(r"\d+\s*[аa]?(?![\d])", s.replace(" ", "")))
    nums = tuple(x.replace("a", "а") for x in nums)
    core = re.sub(r"[№\d.,()\-–/]+", " ", s)
    core = re.sub(STOP, " ", core)
    core = re.sub(r"\s+", " ", core).strip()
    if re.fullmatch(r"[ivxlc]+", core):          # римские номера (участки Балхаша) — как номер
        return kind, nums + (core,), ""
    return kind, nums, core


def similar(a: str, b: str) -> bool:
    if not a or not b:
        return a == b
    a, b = a.replace(" ", ""), b.replace(" ", "")
    if a == b or difflib.SequenceMatcher(None, a, b).ratio() >= 0.8:
        return True
    return min(len(a), len(b)) >= 5 and (a.startswith(b) or b.startswith(a))      # «камышзав» — «камышзавод»


def group_places(places: list[dict]) -> list[list[dict]]:
    """places: [{id, name, rows}] одного водоёма → группы (объединение по связности)."""
    feats = [features(p["name"]) for p in places]
    parent = list(range(len(places)))

    def find(i):
        while parent[i] != i:
            parent[i] = parent[parent[i]]
            i = parent[i]
        return i
    for i in range(len(places)):
        for j in range(i + 1, len(places)):
            (ki, ni, ci), (kj, nj, cj) = feats[i], feats[j]
            same_kind = ki == kj or {ki, kj} == {"sci", "st"}
            if not same_kind or ni != nj:
                continue
            if ki in ("pump", "prom", "sci", "st") and ni:       # станции и участки — по номеру
                if ci and cj and not similar(ci, cj):
                    continue
            elif not similar(ci, cj):
                continue
            parent[find(i)] = find(j)
    groups = defaultdict(list)
    for i, p in enumerate(places):
        groups[find(i)].append((p, feats[i]))
    out = []
    for g in groups.values():
        g.sort(key=lambda x: -x[0]["rows"])
        out.append(g)
    out.sort(key=lambda g: (-len(g), g[0][0]["name"]))
    return out


def proposal(g) -> tuple[str, str]:
    kinds = {f[0] for _, f in g}
    nums = g[0][1][1]
    note = []
    name = g[0][0]["name"]
    k = g[0][1][0]
    if k in ("pump", "prom", "sci", "st") and len(nums) == 1:
        name = dict((kk, t) for kk, _, t in KINDS)["sci" if k == "st" else k].format(n=nums[0].upper())
    if "st" in kinds:
        note.append("«станция» без «науч.» — проверьте, что это научная станция")
    cores = {f[2] for _, f in g}
    if len(cores) > 1 and k not in ("pump", "prom", "sci", "st"):
        note.append("названия различаются написанием — проверьте")
    return name, "; ".join(note)


def client(a) -> httpx.Client:
    c = httpx.Client(base_url=a.url.rstrip("/") + "/api", timeout=120)
    pw = a.password or getpass.getpass(f"Пароль {a.login}: ")
    r = c.post("/users/auth/login", json={"login": a.login, "password": pw})
    if r.status_code != 200:
        sys.exit(f"Вход не выполнен: {r.text}")
    c.headers["Authorization"] = "Bearer " + r.json()["access_token"]
    return c


def row_counts(c) -> dict[tuple[str, str], int]:
    cnt = Counter()
    for d in c.get("/data-admin/datasets", params={"dataset_type": "catches"}).raise_for_status().json():
        t = c.get(f"/data-admin/datasets/{d['id']}/data").raise_for_status().json().get("table") or {}
        for r in t.get("rows") or []:
            if r.get("place"):
                cnt[(d["water_body"], str(r["place"]).strip())] += 1
    return cnt


def suggest(a):
    c = client(a)
    wbs = {r["values"]["code"]: r["values"]["name_ru"]
           for r in c.get("/data-model/entities/water_bodies/records").raise_for_status().json()}
    recs = c.get("/data-model/entities/fishing_places/records").raise_for_status().json()
    cnt = row_counts(c)
    by_wb = defaultdict(list)
    for r in recs:
        v = r["values"]
        aliases = [x for x in str(v.get("aliases") or "").splitlines() if x.strip()]
        rows = cnt.get((v["water_body"], v["name_ru"]), 0) + sum(cnt.get((v["water_body"], x), 0) for x in aliases)
        by_wb[v["water_body"]].append({"id": r["id"], "name": v["name_ru"], "aliases": aliases, "rows": rows})
    wb = Workbook()
    ws = wb.active
    ws.title = "Объединение"
    head = ["Водоём", "Группа", "Название места", "Написание в справочнике", "Другие написания", "Строк уловов",
            "Решение", "Примечание", "id записи"]
    ws.append(head)
    single = wb.create_sheet("Без пары")
    single.append(["Водоём", "Название места", "Другие написания", "Строк уловов", "id записи"])
    fill = PatternFill("solid", fgColor="DCE8F2")
    for cell in list(ws[1]) + list(single[1]):
        cell.font = Font(bold=True)
        cell.fill = fill
    gno = 0
    shade = PatternFill("solid", fgColor="F3F6F9")
    for wbc in sorted(by_wb, key=lambda x: wbs.get(x, x)):
        for g in group_places(by_wb[wbc]):
            if len(g) == 1:
                p = g[0][0]
                single.append([wbs.get(wbc, wbc), p["name"], "; ".join(p["aliases"]), p["rows"], p["id"]])
                continue
            gno += 1
            name, note = proposal(g)
            for i, (p, _) in enumerate(g):
                ws.append([wbs.get(wbc, wbc), gno, name if i == 0 else "", p["name"], "; ".join(p["aliases"]),
                           p["rows"], "объединить" if i == 0 else "", note if i == 0 else "", p["id"]])
                if gno % 2 == 0:
                    for cell in ws[ws.max_row]:
                        cell.fill = shade
    dv = DataValidation(type="list", formula1='"объединить,отдельно"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"G2:G{max(ws.max_row, 2)}")
    for sh, widths in ((ws, [16, 8, 30, 30, 26, 12, 13, 46, 30]), (single, [16, 34, 30, 12, 30])):
        for i, w in enumerate(widths):
            sh.column_dimensions[chr(65 + i)].width = w
        sh.freeze_panes = "A2"
        for row in sh.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
    info = wb.create_sheet("Как заполнить")
    for line in [
        "Места лова — предложение объединить разные написания одного места (FishStock).",
        "Лист «Объединение»: одна группа — одно место. В первой строке группы:",
        "  «Название места» — как место будет называться (можно исправить);",
        "  «Решение» — «объединить» или «отдельно» (тогда все написания группы останутся отдельными местами).",
        "Если в группу попало лишнее написание, удалите его строку — оно останется отдельным местом.",
        "Лист «Без пары» — места, для которых похожих написаний не найдено; их можно не трогать.",
        "Загруженные данные не меняются; при следующих загрузках и правках написания приводятся к названию места.",
    ]:
        info.append([line])
    info.column_dimensions["A"].width = 110
    wb.save(a.out)
    print(f"{a.out}: групп {gno}, мест без пары {single.max_row - 1}")


def load_groups(path: str, wb_by_name: dict[str, str]) -> list[tuple[str, str, list[str]]]:
    """[(код водоёма, название места, все написания)] из таблицы ихтиолога (.xlsx) или из
    services/data_admin/app/seed/place_groups.json (подтверждённые группы поставки)."""
    out = []
    if path.endswith(".json"):
        for wb, gs in json.loads(open(path, encoding="utf-8").read()).items():
            if wb.startswith("_"):
                continue
            out += [(wb, g["name"], [g["name"], *g["spellings"]]) for g in gs]
        return out
    ws = load_workbook(path)["Объединение"]
    groups = defaultdict(list)
    for r in ws.iter_rows(min_row=2, values_only=True):
        if r[1] is not None and r[3]:
            groups[int(r[1])].append(r)
    for _, rows in sorted(groups.items()):
        head = rows[0]
        if str(head[6] or "").strip().lower() != "объединить" or len(rows) < 2:
            continue
        spell = []
        for r in rows:
            spell += [r[3], *[x.strip() for x in str(r[4] or "").split(";") if x.strip()]]
        out.append((wb_by_name.get(head[0], head[0]), str(head[2] or head[3]).strip(), spell))
    return out


def apply(a):
    c = client(a)
    wbs = {r["values"]["name_ru"]: r["values"]["code"]
           for r in c.get("/data-model/entities/water_bodies/records").raise_for_status().json()}
    recs = c.get("/data-model/entities/fishing_places/records").raise_for_status().json()
    done = 0
    for wb, name, spellings in load_groups(a.file, wbs):
        keys = {key(x) for x in spellings} | {key(name)}
        hit = []
        for r in recs:
            v = r["values"]
            if r.get("deleted") or v.get("water_body") != wb:
                continue
            own = [v["name_ru"], *[x for x in str(v.get("aliases") or "").splitlines() if x.strip()]]
            if {key(x) for x in own} & keys:
                hit.append((r, own))
        if not hit:
            continue
        hit.sort(key=lambda h: key(h[0]["values"]["name_ru"]) != key(name))     # запись с названием группы — первой
        main = hit[0][0]
        spell = [x for _, own in hit for x in own] + list(spellings)
        seen, aliases = {key(name)}, []
        for x in (s.strip() for s in spell):            # одно написание на ключ (без повторов регистра и пробелов)
            if x and key(x) not in seen:
                seen.add(key(x))
                aliases.append(x)
        vals = dict(main["values"], name_ru=name, aliases="\n".join(aliases))
        if len(hit) == 1 and vals == main["values"]:
            continue
        print(f"{wb}: «{name}» ← {', '.join(aliases)}")
        if a.dry_run:
            continue
        c.put(f"/data-model/entities/fishing_places/records/{main['id']}",
              json={"values": vals, "comment": "объединение написаний (подтверждено ихтиологом)"}).raise_for_status()
        for r, _ in hit[1:]:
            c.delete(f"/data-model/entities/fishing_places/records/{r['id']}").raise_for_status()
            r["deleted"] = True
        main["values"] = vals
        done += 1
    print(f"объединено групп: {done}" + (" (проверка без изменений)" if a.dry_run else ""))


def key(s) -> str:
    """Как в сервисе данных: без регистра, пробелов, точек, дефисов и кавычек; ё = е."""
    return re.sub(r"[\s.\-–—,«»\"'()]+", "", str(s or "").lower().replace("ё", "е"))


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("command", choices=["suggest", "apply"])
    ap.add_argument("--url", default="http://127.0.0.1:8080")
    ap.add_argument("--login", default="admin")
    ap.add_argument("--password", help="если не указан — будет запрошен")
    ap.add_argument("--out", default="places_groups.xlsx")
    ap.add_argument("--file")
    ap.add_argument("--dry-run", action="store_true")
    a = ap.parse_args()
    if a.command == "suggest":
        suggest(a)
    else:
        if not a.file:
            sys.exit("Укажите --file с заполненной таблицей")
        apply(a)


if __name__ == "__main__":
    main()
