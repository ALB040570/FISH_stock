"""
Таблицы наборов данных: Excel-шаблон, чтение Excel, проверка по типу набора данных, сравнение версий.

Формат хранения версии (data.json):
    {"columns": [коды столбцов], "species": [коды видов], "rows": [{код: значение, ...}]}
Значения столбцов видов хранятся под кодом вида. Пустое значение (null) — пропуск, не ноль.
"""
from __future__ import annotations

import io
import re
from collections import Counter

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from fishcommon.schema import coerce, is_missing

DATA_SHEET = "Данные"
INFO_SHEET = "Описание"
HEADER_FILL = PatternFill("solid", fgColor="DCE8F2")
REQ_FILL = PatternFill("solid", fgColor="F6D7A7")
SPECIES_FILL = PatternFill("solid", fgColor="DDEFD9")


def _norm(s) -> str:
    s = str(s or "").strip().lower().replace("ё", "е")
    s = re.sub(r"\s*\*$", "", s)
    s = re.sub(r"\s*\([^)]*\)\s*$", "", s)   # «Длина, м (обяз.)» → «длина, м»
    return re.sub(r"\s+", " ", s)


def header_title(col: dict, lang: str = "ru") -> str:
    t = col.get(f"title_{lang}") or col.get("title_ru") or col["code"]
    return t + (" *" if col.get("required") else "")


def species_value_field(dtype: dict, sp: dict) -> dict:
    sc = dtype.get("species_columns") or {}
    return {"code": sp["code"], "title_ru": sp.get("name_ru") or sp["code"], "type": sc.get("value_type", "number"),
            "required": False, "min": sc.get("min"), "unit": sc.get("unit")}


# ------------------------------------------------------------------ шаблон
def build_template(dtype: dict, species: list[dict], lookups: dict[str, list[str]], lang: str = "ru",
                   rows: list[dict] | None = None, ref_display: dict[str, dict] | None = None) -> bytes:
    """Excel-шаблон: лист «Данные» (заголовки, выпадающие списки) и лист «Описание»."""
    wb = Workbook()
    ws = wb.active
    ws.title = DATA_SHEET
    cols = dtype["columns"]
    sc = dtype.get("species_columns") or {}
    sp_list = species if sc.get("enabled") else []
    headers = [header_title(c, lang) for c in cols] + [
        (s.get(f"name_{lang}") or s.get("name_ru") or s["code"]) for s in sp_list]
    ws.append(headers)
    for i, c in enumerate(cols, start=1):
        cell = ws.cell(row=1, column=i)
        cell.fill = REQ_FILL if c.get("required") else HEADER_FILL
        cell.font = Font(bold=True)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        ws.column_dimensions[get_column_letter(i)].width = max(12, min(30, len(headers[i - 1]) + 2))
    for j in range(len(sp_list)):
        cell = ws.cell(row=1, column=len(cols) + j + 1)
        cell.fill = SPECIES_FILL
        cell.font = Font(bold=True)
        ws.column_dimensions[get_column_letter(len(cols) + j + 1)].width = 11
    ws.freeze_panes = "A2"
    ws.row_dimensions[1].height = 45

    # выпадающие списки для перечислений и ссылок
    for i, c in enumerate(cols, start=1):
        options = None
        if c["type"] == "enum":
            options = [e.get(f"title_{lang}") or e.get("title_ru") or e["value"] for e in c.get("enum", [])]
        elif c["type"] == "ref":
            options = lookups.get(c.get("ref_entity"), [])
        elif c["type"] == "boolean":
            options = ["да", "нет"] if lang == "ru" else ["иә", "жоқ"]
        if options and sum(len(o) + 1 for o in options) < 250:
            dv = DataValidation(type="list", formula1='"' + ",".join(o.replace('"', "'") for o in options) + '"',
                                allow_blank=True, showErrorMessage=True,
                                error="Выберите значение из списка", errorTitle="FishStocke")
            dv.add(f"{get_column_letter(i)}2:{get_column_letter(i)}5000")
            ws.add_data_validation(dv)

    for r in rows or []:
        line = []
        for c in cols:
            v = r.get(c["code"])
            if c["type"] == "enum":
                v = next((e.get(f"title_{lang}") or e.get("title_ru") for e in c.get("enum", []) if e["value"] == v), v)
            elif c["type"] == "ref" and v is not None and ref_display:
                v = (ref_display.get(c.get("ref_entity")) or {}).get(v, v)
            elif c["type"] == "boolean" and v is not None:
                v = ("да" if v else "нет") if lang == "ru" else ("иә" if v else "жоқ")
            elif c["type"] == "date" and v:
                y, m, d = str(v).split("-")
                v = f"{d}.{m}.{y}"
            elif c["type"] == "geo_point" and v:
                v = f"{v[0]}, {v[1]}"
            line.append(v)
        line += [r.get(s["code"]) for s in sp_list]
        ws.append(line)

    info = wb.create_sheet(INFO_SHEET)
    info.append([dtype.get(f"title_{lang}") or dtype["title_ru"]])
    info["A1"].font = Font(bold=True, size=13)
    info.append([dtype.get("description", "")])
    info.append([])
    info.append(["Столбец", "Код", "Тип", "Единицы", "Обязательный", "Допустимые значения / пределы", "Подсказка"])
    for c in info[4]:
        c.font = Font(bold=True)
        c.fill = HEADER_FILL
    from fishcommon.schema import TYPE_TITLES
    for c in cols:
        allowed = ""
        if c["type"] == "enum":
            allowed = ", ".join(e.get("title_ru") or e["value"] for e in c.get("enum", []))
        elif c["type"] == "ref":
            allowed = f"из справочника «{c.get('ref_entity')}»: " + ", ".join(lookups.get(c.get("ref_entity"), [])[:30])
        elif c["type"] in ("integer", "number"):
            parts = []
            if c.get("min") is not None:
                parts.append(f"не меньше {c['min']}")
            if c.get("max") is not None:
                parts.append(f"не больше {c['max']}")
            allowed = ", ".join(parts)
        elif c["type"] == "date":
            allowed = "ДД.ММ.ГГГГ"
        info.append([c.get("title_ru"), c["code"], TYPE_TITLES.get(c["type"], c["type"]), c.get("unit", ""),
                     "да" if c.get("required") else "нет", allowed, c.get("hint", "")])
    if sc.get("enabled"):
        info.append([])
        info.append([f"Столбцы видов (зелёные): {sc.get('title_ru', '')}", "", TYPE_TITLES.get(sc.get("value_type"), ""),
                     sc.get("unit", ""), "нет", f"не меньше {sc.get('min', 0)}", sc.get("missing", "")])
        info.append(["Лишние столбцы видов можно удалить; новые виды сначала добавьте в справочник «Виды рыб»."])
    info.append([])
    info.append(["Заголовки на листе «Данные» не переименовывайте. Звёздочка * — обязательный столбец."])
    for col, w in zip("ABCDEFG", (34, 16, 16, 10, 12, 50, 50)):
        info.column_dimensions[col].width = w
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ------------------------------------------------------------------ чтение Excel
class ParseError(Exception):
    def __init__(self, errors: list[dict]):
        super().__init__("parse")
        self.errors = errors


def read_excel(content: bytes, dtype: dict, species: list[dict]) -> tuple[list[str], list[tuple[int, dict]]]:
    """Возвращает (коды видов в файле, [(номер строки Excel, {код: сырое значение})])."""
    try:
        wb = load_workbook(io.BytesIO(content), data_only=True, read_only=True)
    except Exception:  # noqa: BLE001
        raise ParseError([{"row": None, "column": None,
                           "message": "Файл не читается как Excel (.xlsx). Сохраните его в формате «Книга Excel»."}])
    ws = wb[DATA_SHEET] if DATA_SHEET in wb.sheetnames else wb.worksheets[0]
    it = ws.iter_rows(values_only=True)
    try:
        header = next(it)
    except StopIteration:
        raise ParseError([{"row": 1, "column": None, "message": "Лист пуст: нет строки заголовков"}])

    cols = dtype["columns"]
    by_name: dict[str, str] = {}
    for c in cols:
        for cand in (c["code"], c.get("title_ru"), c.get("title_kk")):
            if cand:
                by_name[_norm(cand)] = c["code"]
    sc = dtype.get("species_columns") or {}
    sp_by_name: dict[str, str] = {}
    if sc.get("enabled"):
        for s in species:
            for cand in (s["code"], s.get("name_ru"), s.get("name_kk"), s.get("latin")):
                if cand:
                    sp_by_name[_norm(cand)] = s["code"]

    mapping: dict[int, str] = {}
    errors = []
    found_species: list[str] = []
    seen = Counter()
    for idx, h in enumerate(header):
        if h is None or str(h).strip() == "":
            continue
        key = _norm(h)
        code = by_name.get(key)
        if code is None and key in sp_by_name:
            code = sp_by_name[key]
            found_species.append(code)
        if code is None:
            msg = (f"Столбец «{h}»: такого столбца нет в типе «{dtype['title_ru']}»"
                   + ("; если это вид рыбы — добавьте его в справочник «Виды рыб»" if sc.get("enabled") else ""))
            errors.append({"row": 1, "column": get_column_letter(idx + 1), "message": msg})
            continue
        seen[code] += 1
        if seen[code] > 1:
            errors.append({"row": 1, "column": get_column_letter(idx + 1), "message": f"Столбец «{h}» повторяется"})
        mapping[idx] = code
    missing = [c for c in cols if c.get("required") and c["code"] not in mapping.values()]
    for c in missing:
        errors.append({"row": 1, "column": None, "message": f"Нет обязательного столбца «{c['title_ru']}»"})
    if errors:
        raise ParseError(errors)

    rows = []
    for rn, values in enumerate(it, start=2):
        if values is None or all(is_missing(v) for v in values):
            continue
        rec = {}
        for idx, code in mapping.items():
            rec[code] = values[idx] if idx < len(values) else None
        rows.append((rn, rec))
    return found_species, rows


# ------------------------------------------------------------------ проверка
def validate_rows(dtype: dict, species_codes: list[str], species: list[dict], raw_rows: list[tuple[int, dict]],
                  ref_maps: dict[str, dict]) -> tuple[list[dict], list[dict], list[dict]]:
    """Проверка строк. Возвращает (чистые строки, ошибки, предупреждения)."""
    cols = dtype["columns"]
    sp_by_code = {s["code"]: s for s in species}
    sp_fields = [species_value_field(dtype, sp_by_code.get(c, {"code": c})) for c in species_codes]
    clean, errors, warnings = [], [], []
    letters = {}
    for i, c in enumerate(cols):
        letters[c["code"]] = c.get("title_ru")
    for rn, rec in raw_rows:
        out = {}
        for c in cols:
            v, err = coerce(c, rec.get(c["code"]), ref_maps.get(c.get("ref_entity") or ""))
            if err:
                errors.append({"row": rn, "column": c.get("title_ru"), "field": c["code"], "message": err})
            out[c["code"]] = v
        for f in sp_fields:
            v, err = coerce(f, rec.get(f["code"]), None)
            if err:
                errors.append({"row": rn, "column": f["title_ru"], "field": f["code"], "message": err})
            out[f["code"]] = v
        for rule in dtype.get("rules") or []:
            a, b = out.get(rule["left"]), out.get(rule["right"])
            if a is not None and b is not None and ((rule["type"] == "le" and a > b) or (rule["type"] == "lt" and a >= b)):
                errors.append({"row": rn, "column": None, "field": rule["left"], "message": rule.get("message")})
        clean.append({"_row": rn, **out})
    keys = dtype.get("key_fields") or []
    if keys:
        cnt = Counter(tuple(r.get(k) for k in keys) for r in clean)
        for r in clean:
            k = tuple(r.get(x) for x in keys)
            if cnt[k] > 1 and None not in k:
                names = ", ".join(f"{letters.get(x, x)} = {v}" for x, v in zip(keys, k))
                errors.append({"row": r["_row"], "column": None, "field": keys[0],
                               "message": f"Повтор строки с тем же ключом ({names})"})
    pf = dtype.get("period_field")
    if pf and clean:
        years = sorted({r.get(pf) for r in clean if r.get(pf) is not None})
        if "date" in {c["code"] for c in cols}:
            for r in clean:
                d, y = r.get("date"), r.get(pf)
                if d and y and int(d[:4]) != int(y):
                    warnings.append({"row": r["_row"], "column": "Дата", "field": "date",
                                     "message": f"Год даты {d[:4]} не совпадает с годом {y}"})
    for r in clean:
        r.pop("_row", None)
    return clean, errors, warnings


def merge_rows(dtype: dict, old: dict | None, new_rows: list[dict], new_species: list[str], mode: str) -> dict:
    """mode: replace — заменить весь набор; merge — заменить периоды, присутствующие в файле."""
    cols = [c["code"] for c in dtype["columns"]]
    if mode == "replace" or not old:
        species = list(new_species)
        return {"columns": cols, "species": species, "rows": new_rows}
    pf = dtype.get("period_field")
    if not pf:
        keys = dtype.get("key_fields") or []
        new_keys = {tuple(r.get(k) for k in keys) for r in new_rows}
        kept = [r for r in old["rows"] if tuple(r.get(k) for k in keys) not in new_keys]
    else:
        periods = {r.get(pf) for r in new_rows}
        kept = [r for r in old["rows"] if r.get(pf) not in periods]
    species = list(dict.fromkeys(list(old.get("species", [])) + list(new_species)))
    rows = kept + new_rows
    for r in rows:
        for s in species:
            r.setdefault(s, None)
    if pf:
        rows.sort(key=lambda r: (r.get(pf) is None, r.get(pf) or 0))
    return {"columns": cols, "species": species, "rows": rows}


def summarize(dtype: dict, table: dict) -> dict:
    pf = dtype.get("period_field")
    rows = table["rows"]
    years = sorted({r.get(pf) for r in rows if r.get(pf) is not None}) if pf else []
    return {"rows": len(rows), "species": table.get("species", []),
            "period_from": years[0] if years else None, "period_to": years[-1] if years else None,
            "years": years}


def diff_tables(dtype: dict, a: dict, b: dict) -> dict:
    """Сравнение двух версий. По ключевым полям — изменённые строки; иначе — добавленные/удалённые."""
    keys = dtype.get("key_fields") or []
    fields = a.get("columns", []) + sorted(set(a.get("species", [])) | set(b.get("species", [])))

    def norm(r):
        return tuple((f, r.get(f)) for f in fields)

    if keys:
        ka = {tuple(r.get(k) for k in keys): r for r in a["rows"]}
        kb = {tuple(r.get(k) for k in keys): r for r in b["rows"]}
        added = [kb[k] for k in kb if k not in ka]
        removed = [ka[k] for k in ka if k not in kb]
        changed = []
        for k in ka.keys() & kb.keys():
            ra, rb = ka[k], kb[k]
            cells = [{"field": f, "old": ra.get(f), "new": rb.get(f)} for f in fields if ra.get(f) != rb.get(f)]
            if cells:
                changed.append({"key": dict(zip(keys, k)), "cells": cells})
        return {"keys": keys, "added": added, "removed": removed, "changed": changed}
    ca, cb = Counter(norm(r) for r in a["rows"]), Counter(norm(r) for r in b["rows"])
    added = [dict(t) for t, n in (cb - ca).items() for _ in range(n)]
    removed = [dict(t) for t, n in (ca - cb).items() for _ in range(n)]
    return {"keys": [], "added": added, "removed": removed, "changed": []}
