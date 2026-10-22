"""
Отчёт по запуску расчёта ПДУ (Word и PDF). Структура предложена разработчиком (согласовано 18.10):

1. Сведения о расчёте: водоём, расчёт, автор, даты, версии алгоритмов, версии наборов данных.
2. Итог: ПДУ по видам (прогноз F = k·M на первый прогнозный год с 95 % интервалом; ПДУ НПЦ последнего года).
3. Запас: окончательный запас по уловам и численность по байесовской модели (последний год истории).
4. ПДУ НПЦ по возрастам (последний год).
5. Диагностика модели и примечания; ограничения.
"""
from __future__ import annotations

import io
from pathlib import Path

FONT_DIR = Path(__file__).parent / "fonts"
STEP_TITLES = {"seine_catchability": "Уловистость невода", "stock_by_catch": "Уловистость сети и запас по уловам",
               "pdu_by_age": "ПДУ НПЦ по возрастам", "bayes_cohort": "Байесовская когортная модель",
               "forecast_pdu": "Прогноз и ПДУ"}


def _fmt(v, digits=0):
    if v is None:
        return "—"
    try:
        f = float(v)
    except (TypeError, ValueError):
        return str(v)
    if digits == 0:
        return f"{f:,.0f}".replace(",", " ")
    return f"{f:,.{digits}f}".replace(",", " ").replace(".", ",")


def _sections(run: dict, results: dict, wb: dict) -> list[tuple[str, list]]:
    """Содержимое отчёта: [(заголовок, [абзац | ("table", заголовки, строки)])]."""
    out = []
    info = [
        ("table", ["Показатель", "Значение"], [
            ["Водоём", wb.get("name_ru", run["water_body"])],
            ["Расчёт", run.get("calculation_title", "")],
            ["Идентификатор запуска", run["id"]],
            ["Автор", run.get("author_name") or run.get("author") or ""],
            ["Начат / завершён (UTC)", f"{run.get('started_at', '—')} / {run.get('finished_at', '—')}"],
            ["Шаги (алгоритм, версия)", "; ".join(f"{STEP_TITLES.get(s['algorithm'], s['algorithm'])} {s['version']}"
                                                   for s in run["steps"])],
            ["Наборы данных (тип: версия)", "; ".join(f"{k}: v{v['version']}" for k, v in sorted(run["datasets"].items()))],
        ])]
    out.append(("1. Сведения о расчёте", info))

    fc = results.get("forecast_pdu") or {}
    rows = (fc.get("tables", {}).get("forecast_summary") or {}).get("rows", [])
    if rows:
        y = rows[0].get("year_first_fc")
        out.append((f"2. ПДУ на {y} год", [
            "ПДУ рассчитан по сценарию F = k·M (уравнение Баранова) по 500 выборкам байесовской модели; "
            "приведены медиана и 95 % интервал. Для сравнения — ПДУ НПЦ последнего года истории.",
            ("table", ["Вид", "ПДУ, т (медиана)", "95 % интервал, т", "F", "ПДУ НПЦ, т", "Вероятность снижения запаса"],
             [[r["species"], _fmt(r["pdu_next_t"], 1), f"{_fmt(r['pdu_next_t_q025'], 1)} – {_fmt(r['pdu_next_t_q975'], 1)}",
               _fmt(r.get("F_next"), 3), _fmt(r.get("npc_last_t"), 1), _fmt(r.get("p_decline_baranov"), 2)] for r in rows]),
            f"Параметры прогноза: k = {_fmt(fc.get('payload', {}).get('k'), 2)}, "
            f"горизонт — {fc.get('payload', {}).get('horizon')} лет."]))
    else:
        out.append(("2. ПДУ", ["В запуске нет шага прогноза — ПДУ по сценариям не рассчитан."]))

    st = results.get("stock_by_catch") or {}
    by = (st.get("tables", {}).get("yearly_stock") or {}).get("rows", [])
    if by:
        last = max(r["year"] for r in by)
        tab = [[r["species"], _fmt(r["stock"])] for r in by
               if r["year"] == last and r["gear"] not in ("сеть", "невод") and r["variant"] == "with_ages"
               and r["method"] == "on_stock"]
        par = [f"Окончательный запас по уловам (сеть и невод, уловистость по возрастам) за {last} г., шт."]
        bc = results.get("bayes_cohort") or {}
        btab = (bc.get("tables", {}).get("bayes_totals") or {}).get("rows", [])
        bmap = {r["species"]: r for r in btab if r["year"] == last}
        out.append(("3. Запас", par + [("table", ["Вид", "Запас по уловам, шт", "Байесовская модель, медиана, шт",
                                                  "95 % интервал, шт"],
                                        [[s, v, _fmt(bmap.get(s, {}).get("total_median")),
                                          f"{_fmt(bmap.get(s, {}).get('total_q025'))} – {_fmt(bmap.get(s, {}).get('total_q975'))}"
                                          if s in bmap else "—"] for s, v in tab])]))

    pd_ = results.get("pdu_by_age") or {}
    prow = (pd_.get("tables", {}).get("pdu_by_age") or {}).get("rows", [])
    if prow:
        last = max(r["year"] for r in prow)
        tab = [[r["species"], r["age_title"], _fmt(r["stock"]), _fmt(r["weight_g"], 1), _fmt(r["pdu_pcs"]), _fmt(r["pdu_t"], 2)]
               for r in prow if r["year"] == last and (r["pdu_pcs"] or 0) > 0]
        out.append((f"4. ПДУ НПЦ по возрастам, {last} г.", [
            "Квота НПЦ распределена между половозрелыми возрастами пропорционально биомассе окончательного запаса "
            "с нагрузкой j/m1; за итерацию из возраста изымается не более 75 % остатка.",
            ("table", ["Вид", "Возраст", "Запас, шт", "Ср. вес, г", "ПДУ, шт", "ПДУ, т"], tab)]))

    notes = []
    for code, res in results.items():
        for n in (res or {}).get("notes", []):
            notes.append(f"{STEP_TITLES.get(code, code)}: {n}")
        for k, m in ((res or {}).get("metrics") or {}).items():
            if k.startswith(("rhat_", "divergences_")) and m.get("value") is not None:
                notes.append(f"{m['title']}: {_fmt(m['value'], 3 if k.startswith('rhat') else 0)}")
    notes += [w for w in run.get("warnings", [])]
    notes.append("Методика расчёта: «Многовидовые когортные математические модели оценки запасов рыбных ресурсов "
                 "с применением байесовского вывода» (Габбасов М. и др.).")
    out.append(("5. Диагностика и примечания", notes))
    return out


def docx_report(run: dict, results: dict, wb: dict) -> bytes:
    from docx import Document
    from docx.shared import Pt
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = "Times New Roman"
    st.font.size = Pt(11)
    doc.add_heading(f"Обоснование ПДУ: {wb.get('name_ru', run['water_body'])}", 0)
    doc.add_paragraph("Отчёт сформирован приложением FishStock.")
    for title, items in _sections(run, results, wb):
        doc.add_heading(title, 1)
        for it in items:
            if isinstance(it, tuple):
                _, head, rows = it
                t = doc.add_table(rows=1, cols=len(head))
                t.style = "Light Grid Accent 1"
                for c, h in zip(t.rows[0].cells, head):
                    c.text = str(h)
                for r in rows:
                    cells = t.add_row().cells
                    for c, v in zip(cells, r):
                        c.text = str(v)
            else:
                doc.add_paragraph(str(it), style="List Bullet" if title.startswith("5.") else None)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def pdf_report(run: dict, results: dict, wb: dict) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    if "DejaVu" not in pdfmetrics.getRegisteredFontNames():
        pdfmetrics.registerFont(TTFont("DejaVu", str(FONT_DIR / "DejaVuSans.ttf")))
        pdfmetrics.registerFont(TTFont("DejaVu-Bold", str(FONT_DIR / "DejaVuSans-Bold.ttf")))
    body = ParagraphStyle("b", fontName="DejaVu", fontSize=9, leading=12)
    h0 = ParagraphStyle("h0", fontName="DejaVu-Bold", fontSize=15, leading=19, spaceAfter=6)
    h1 = ParagraphStyle("h1", fontName="DejaVu-Bold", fontSize=11.5, leading=15, spaceBefore=10, spaceAfter=4)
    cell = ParagraphStyle("c", fontName="DejaVu", fontSize=8, leading=10)
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=landscape(A4), leftMargin=15 * mm, rightMargin=15 * mm,
                            topMargin=12 * mm, bottomMargin=12 * mm, title="Обоснование ПДУ")
    story = [Paragraph(f"Обоснование ПДУ: {wb.get('name_ru', run['water_body'])}", h0),
             Paragraph("Отчёт сформирован приложением FishStock.", body)]
    for title, items in _sections(run, results, wb):
        story.append(Paragraph(title, h1))
        for it in items:
            if isinstance(it, tuple):
                _, head, rows = it
                data = [[Paragraph(f"<b>{h}</b>", cell) for h in head]] + [[Paragraph(str(v), cell) for v in r] for r in rows]
                t = Table(data, repeatRows=1)
                t.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#9aa9b8")),
                                       ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#dce8f2")),
                                       ("VALIGN", (0, 0), (-1, -1), "TOP")]))
                story += [t, Spacer(1, 4)]
            else:
                story.append(Paragraph(("• " if title.startswith("5.") else "") + str(it), body))
    doc.build(story)
    return buf.getvalue()
