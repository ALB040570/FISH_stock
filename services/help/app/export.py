"""
Выгрузка справки в Word и PDF. Markdown разбирается markdown-it (CommonMark + таблицы + формулы $…$ и $$…$$)
в простые блоки: заголовки, абзацы, списки, таблицы, код, цитаты (врезки «Важно», «[Снимок экрана: …]»),
рисунки, формулы. Формулы (подмножество LaTeX) рисуются matplotlib (mathtext) в PNG.
Рисунки статей — ![подпись](help-image:имя) (загружены в приложении) и ![подпись](bundled:имя) (поставка).
"""
from __future__ import annotations

import io
import re
from datetime import date
from pathlib import Path
from typing import Callable

from markdown_it import MarkdownIt
from mdit_py_plugins.dollarmath import dollarmath_plugin

FONT_DIR = Path(__file__).parent / "fonts"
T = {
    "ru": {"guide": "Руководство пользователя", "app": "FishStocke — определение предельно допустимого улова рыб "
           "в водоёмах Казахстана", "toc": "Содержание", "date": "Дата выгрузки", "page": "стр.",
           "noimg": "рисунок не найден"},
    "kk": {"guide": "Пайдаланушы нұсқаулығы", "app": "FishStocke — Қазақстан су айдындарында балық аулаудың шекті "
           "рұқсат етілген көлемін анықтау", "toc": "Мазмұны", "date": "Түсірілген күні", "page": "бет",
           "noimg": "сурет табылмады"},
}
MD = MarkdownIt("commonmark").enable("table").use(dollarmath_plugin, allow_space=True, double_inline=True)
_MATH_CACHE: dict[tuple[str, int], tuple[bytes, int, int]] = {}


def math_png(tex: str, size: int = 11, dpi: int = 300) -> tuple[bytes, int, int]:
    """Формула → (PNG, ширина, высота в пикселях при dpi)."""
    key = (tex, size)
    if key not in _MATH_CACHE:
        import matplotlib
        matplotlib.use("Agg")
        from matplotlib import mathtext
        from matplotlib.font_manager import FontProperties
        from PIL import Image as PILImage
        buf = io.BytesIO()
        mathtext.math_to_image(f"${tex.strip()}$", buf, prop=FontProperties(size=size), dpi=dpi, format="png")
        data = buf.getvalue()
        w, h = PILImage.open(io.BytesIO(data)).size
        _MATH_CACHE[key] = (data, w, h)
    return _MATH_CACHE[key]
Run = tuple[str, dict]          # (текст, {"b", "i", "code"})


# ------------------------------------------------------------------ разбор Markdown
def _inline(tok) -> list:
    """Инлайн-токен → список прогонов и рисунков [("run", текст, стиль) | ("img", src, alt)]."""
    out, st = [], {"b": False, "i": False, "code": False}
    for c in tok.children or []:
        t = c.type
        if t == "text":
            out.append(("run", c.content, dict(st)))
        elif t == "code_inline":
            out.append(("run", c.content, {**st, "code": True}))
        elif t in ("softbreak", "hardbreak"):
            out.append(("run", " " if t == "softbreak" else "\n", dict(st)))
        elif t == "strong_open":
            st["b"] = True
        elif t == "strong_close":
            st["b"] = False
        elif t == "em_open":
            st["i"] = True
        elif t == "em_close":
            st["i"] = False
        elif t == "math_inline":
            out.append(("math", c.content, dict(st)))
        elif t == "image":
            alt = "".join(x.content for x in (c.children or []))
            out.append(("img", c.attrs.get("src", ""), alt))
        elif t in ("html_inline",):
            continue
    return out


def parse(md: str) -> list:
    """Markdown → блоки: ("h", n, inl) ("p", inl) ("list", ordered, [[блоки]]) ("table", head, rows)
    ("code", текст) ("quote", [блоки]) ("hr",)."""
    tokens = MD.parse(md)
    pos = 0

    def block_seq(stop: str | None) -> list:
        nonlocal pos
        out = []
        while pos < len(tokens):
            tk = tokens[pos]
            if stop and tk.type == stop:
                pos += 1
                return out
            if tk.type == "heading_open":
                out.append(("h", int(tk.tag[1]), _inline(tokens[pos + 1])))
                pos += 3
            elif tk.type == "paragraph_open":
                out.append(("p", _inline(tokens[pos + 1])))
                pos += 3
            elif tk.type in ("bullet_list_open", "ordered_list_open"):
                ordered = tk.type == "ordered_list_open"
                close = tk.type.replace("_open", "_close")
                pos += 1
                items = []
                while tokens[pos].type != close:
                    if tokens[pos].type == "list_item_open":
                        pos += 1
                        items.append(block_seq("list_item_close"))
                    else:
                        pos += 1
                pos += 1
                out.append(("list", ordered, items))
            elif tk.type == "blockquote_open":
                pos += 1
                out.append(("quote", block_seq("blockquote_close")))
            elif tk.type in ("math_block", "math_block_label"):
                out.append(("mathblock", tk.content.strip()))
                pos += 1
            elif tk.type in ("fence", "code_block"):
                out.append(("code", tk.content.rstrip("\n")))
                pos += 1
            elif tk.type == "hr":
                out.append(("hr",))
                pos += 1
            elif tk.type == "table_open":
                head, rows, cur, in_head = [], [], None, False
                pos += 1
                while tokens[pos].type != "table_close":
                    t = tokens[pos].type
                    if t == "thead_open":
                        in_head = True
                    elif t == "thead_close":
                        in_head = False
                    elif t == "tr_open":
                        cur = []
                    elif t == "tr_close":
                        (head if in_head else rows).append(cur)
                    elif t == "inline":
                        cur.append(_inline(tokens[pos]))
                    pos += 1
                pos += 1
                out.append(("table", head[0] if head else [], rows))
            else:
                pos += 1
        return out

    return block_seq(None)


def _plain(inl) -> str:
    return "".join(x[1] for x in inl if x[0] in ("run", "math"))


def _strip_title(md: str) -> str:
    """Первая строка «# Заголовок» выводится отдельно (заголовок статьи)."""
    lines = md.splitlines()
    if lines and lines[0].startswith("# "):
        lines = lines[1:]
    return "\n".join(lines)


BUNDLED_IMAGES = Path(__file__).parent / "content" / "images"


def _img_file(src: str, image_path: Callable) -> Path | None:
    if src.startswith("help-image:"):
        return image_path(src.split(":", 1)[1])
    if src.startswith("bundled:"):
        p = BUNDLED_IMAGES / Path(src.split(":", 1)[1]).name
        return p if p.exists() else None
    return None


# ------------------------------------------------------------------ Word
def to_docx(title: str, parts: list[dict], lang: str, image_path: Callable) -> bytes:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor

    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(11)
    for s in doc.sections:
        s.left_margin = s.right_margin = Cm(2)

    def runs(p, inl):
        for x in inl:
            if x[0] == "math":
                png, w, h = math_png(x[1])
                p.add_run().add_picture(io.BytesIO(png), height=Cm(h / 300 * 2.54))
                continue
            if x[0] == "img":
                f = _img_file(x[1], image_path)
                if f:
                    p.add_run().add_picture(str(f), width=Cm(16))
                else:
                    r = p.add_run(f"[{x[2] or T[lang]['noimg']}]")
                    r.italic = True
                continue
            r = p.add_run(x[1])
            r.bold, r.italic = x[2]["b"], x[2]["i"]
            if x[2]["code"]:
                r.font.name = "Consolas"
                r.font.size = Pt(10)

    def blocks(bl, level=0, quote=False):
        for b in bl:
            k = b[0]
            if k == "h":
                doc.add_heading(_plain(b[2]), level=min(b[1], 6))
            elif k == "p":
                p = doc.add_paragraph()
                if level:
                    p.paragraph_format.left_indent = Cm(0.8 * level)
                if quote:
                    p.paragraph_format.left_indent = Cm(0.8)
                    for x in b[1]:
                        if x[0] == "run":
                            r = p.add_run(x[1])
                            r.italic = True
                            r.bold = x[2]["b"]
                            r.font.color.rgb = RGBColor(0x44, 0x55, 0x66)
                else:
                    runs(p, b[1])
            elif k == "list":
                style = "List Number" if b[1] else "List Bullet"
                if level:
                    style += " 2"
                for item in b[2]:
                    first = True
                    for sub in item:
                        if sub[0] == "p" and first:
                            p = doc.add_paragraph(style=style)
                            runs(p, sub[1])
                            first = False
                        else:
                            blocks([sub], level + 1)
            elif k == "table":
                ncol = max([len(b[2][0]) if b[2] else 0, len(b[1])])
                t = doc.add_table(rows=0, cols=ncol)
                t.style = "Table Grid"
                if b[1]:
                    cells = t.add_row().cells
                    for i, c in enumerate(b[1]):
                        cells[i].text = ""
                        r = cells[i].paragraphs[0].add_run(_plain(c))
                        r.bold = True
                for row in b[2]:
                    cells = t.add_row().cells
                    for i, c in enumerate(row[:ncol]):
                        cells[i].text = ""
                        runs(cells[i].paragraphs[0], c)
                doc.add_paragraph()
            elif k == "mathblock":
                png, w, h = math_png(b[1], size=12)
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                width = min(w / 300 * 2.54, 16.0)
                p.add_run().add_picture(io.BytesIO(png), width=Cm(width))
            elif k == "code":
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Cm(0.5)
                r = p.add_run(b[1])
                r.font.name = "Consolas"
                r.font.size = Pt(9)
            elif k == "quote":
                blocks(b[1], level, quote=True)
            elif k == "hr":
                doc.add_paragraph("—" * 20).alignment = WD_ALIGN_PARAGRAPH.CENTER

    # титульный лист
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("FishStocke")
    r.bold, r.font.size = True, Pt(28)
    p = doc.add_paragraph(T[lang]["app"])
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(title)
    r.bold, r.font.size = True, Pt(20)
    p = doc.add_paragraph(f"{T[lang]['date']}: {date.today().strftime('%d.%m.%Y')}")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # оглавление (поле Word: обновляется при открытии — F9 / «Обновить поле»)
    doc.add_page_break()
    doc.add_heading(T[lang]["toc"], level=1)
    par = doc.add_paragraph()
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), 'TOC \\o "1-3" \\h \\z \\u')
    par._p.append(fld)
    settings = doc.settings.element
    upd = OxmlElement("w:updateFields")
    upd.set(qn("w:val"), "true")
    settings.append(upd)

    single = len(parts) == 1 and len(parts[0]["articles"]) == 1
    for part in parts:
        doc.add_page_break()
        if not single:
            doc.add_heading(part["title"], level=1)
        for a in part["articles"]:
            doc.add_heading(a["title"], level=1 if single else 2)
            blocks(_shift(parse(_strip_title(a["markdown"])), 0 if single else 1))
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _shift(bl: list, by: int) -> list:
    """Сдвиг уровней заголовков внутри статьи (## → уровень 3 в разделе)."""
    if not by:
        return bl
    out = []
    for b in bl:
        if b[0] == "h":
            out.append(("h", b[1] + by, b[2]))
        else:
            out.append(b)
    return out


# ------------------------------------------------------------------ PDF
def to_pdf(title: str, parts: list[dict], lang: str, image_path: Callable) -> bytes:
    from xml.sax.saxutils import escape

    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.platypus import (Image, KeepTogether, ListFlowable, ListItem, PageBreak, Paragraph, Preformatted,
                                    Spacer, Table, TableStyle)
    from reportlab.platypus.doctemplate import BaseDocTemplate, PageTemplate
    from reportlab.platypus.frames import Frame
    from reportlab.platypus.tableofcontents import TableOfContents

    if "DejaVu" not in pdfmetrics.getRegisteredFontNames():
        pdfmetrics.registerFont(TTFont("DejaVu", str(FONT_DIR / "DejaVuSans.ttf")))
        pdfmetrics.registerFont(TTFont("DejaVu-Bold", str(FONT_DIR / "DejaVuSans-Bold.ttf")))
        from reportlab.lib.fonts import addMapping
        addMapping("DejaVu", 0, 0, "DejaVu")
        addMapping("DejaVu", 1, 0, "DejaVu-Bold")
        addMapping("DejaVu", 0, 1, "DejaVu")
        addMapping("DejaVu", 1, 1, "DejaVu-Bold")
    base = ParagraphStyle("b", fontName="DejaVu", fontSize=10, leading=14, spaceAfter=5)
    H = {1: ParagraphStyle("h1", parent=base, fontName="DejaVu-Bold", fontSize=17, leading=22, spaceBefore=6,
                           spaceAfter=10),
         2: ParagraphStyle("h2", parent=base, fontName="DejaVu-Bold", fontSize=14, leading=18, spaceBefore=10,
                           spaceAfter=6),
         3: ParagraphStyle("h3", parent=base, fontName="DejaVu-Bold", fontSize=12, leading=16, spaceBefore=8,
                           spaceAfter=4),
         4: ParagraphStyle("h4", parent=base, fontName="DejaVu-Bold", fontSize=10.5, leading=14, spaceBefore=6)}
    quote = ParagraphStyle("q", parent=base, textColor=colors.HexColor("#445566"), leftIndent=8 * mm,
                           backColor=colors.HexColor("#f2f5f8"), borderPadding=4)
    code = ParagraphStyle("c", parent=base, fontName="DejaVu", fontSize=8.5, leading=11, leftIndent=5 * mm,
                          backColor=colors.HexColor("#f4f4f4"))
    cell = ParagraphStyle("t", parent=base, fontSize=8.5, leading=11, spaceAfter=0)
    cell_b = ParagraphStyle("tb", parent=cell, fontName="DejaVu-Bold")
    width = A4[0] - 40 * mm

    import tempfile
    tmp = Path(tempfile.mkdtemp(prefix="fishhelp-"))

    def math_file(tex, size=10):
        png, w, h = math_png(tex, size=size)
        f = tmp / f"m{abs(hash((tex, size)))}.png"
        if not f.exists():
            f.write_bytes(png)
        return f, w / 300 * 72, h / 300 * 72

    def markup(inl, italic=False) -> tuple[str, list]:
        s, imgs = "", []
        for x in inl:
            if x[0] == "math":
                f, w, h = math_file(x[1])
                s += f'<img src="{f}" width="{w:.1f}" height="{h:.1f}" valign="middle"/>'
                continue
            if x[0] == "img":
                imgs.append(x)
                continue
            t = escape(x[1]).replace("\n", "<br/>")
            if x[2]["code"]:
                t = f'<font face="DejaVu" color="#7a3e00">{t}</font>'
            if x[2]["b"]:
                t = f"<b>{t}</b>"
            if x[2]["i"] or italic:
                t = f"<i>{t}</i>"
            s += t
        return s, imgs

    def image(x):
        f = _img_file(x[1], image_path)
        if not f:
            return Paragraph(f"<i>[{escape(x[2] or T[lang]['noimg'])}]</i>", quote)
        im = Image(str(f))
        k = min(1.0, width / im.imageWidth)
        im.drawWidth, im.drawHeight = im.imageWidth * k, im.imageHeight * k
        return im

    def flow(bl, style=base) -> list:
        out = []
        for b in bl:
            k = b[0]
            if k == "h":
                lvl = min(b[1], 4)
                p = Paragraph(escape(_plain(b[2])), H[lvl])
                p._toc_level = lvl - 1
                out.append(p)
            elif k == "p":
                s, imgs = markup(b[1])
                if s.strip():
                    out.append(Paragraph(s, style))
                out += [image(x) for x in imgs]
            elif k == "list":
                items = []
                for item in b[2]:
                    items.append(ListItem(flow(item, style), leftIndent=12))
                out.append(ListFlowable(items, bulletType="1" if b[1] else "bullet", start="1" if b[1] else None,
                                        bulletFontName="DejaVu", bulletFontSize=9, leftIndent=12))
            elif k == "table":
                data = []
                if b[1]:
                    data.append([Paragraph(markup(c)[0], cell_b) for c in b[1]])
                for row in b[2]:
                    data.append([Paragraph(markup(c)[0], cell) for c in row])
                if data:
                    n = max(len(r) for r in data)
                    data = [r + [""] * (n - len(r)) for r in data]
                    t = Table(data, colWidths=[width / n] * n, repeatRows=1 if b[1] else 0)
                    t.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.4, colors.grey),
                                           ("VALIGN", (0, 0), (-1, -1), "TOP"),
                                           ("BACKGROUND", (0, 0), (-1, 0 if b[1] else -1),
                                            colors.HexColor("#e8eef4") if b[1] else colors.white)]))
                    out += [t, Spacer(1, 6)]
            elif k == "mathblock":
                f, w, h = math_file(b[1], size=13)
                k_ = min(1.0, width / w)
                im = Image(str(f), width=w * k_, height=h * k_)
                im.hAlign = "CENTER"
                out += [Spacer(1, 3), im, Spacer(1, 3)]
            elif k == "code":
                out.append(Preformatted(b[1], code))
            elif k == "quote":
                out.append(KeepTogether(flow(b[1], quote)))
            elif k == "hr":
                out.append(Spacer(1, 8))
        return out

    class Doc(BaseDocTemplate):
        def __init__(self, buf):
            super().__init__(buf, pagesize=A4, leftMargin=20 * mm, rightMargin=20 * mm, topMargin=18 * mm,
                             bottomMargin=18 * mm, title=title, author="FishStocke")
            frame = Frame(self.leftMargin, self.bottomMargin, self.width, self.height, id="f")
            self.addPageTemplates([PageTemplate("p", [frame], onPage=self._footer)])

        def _footer(self, canv, doc):
            canv.saveState()
            canv.setFont("DejaVu", 8)
            canv.setFillColor(colors.grey)
            canv.drawString(20 * mm, 10 * mm, f"FishStocke — {title}")
            canv.drawRightString(A4[0] - 20 * mm, 10 * mm, f"{T[lang]['page']} {doc.page}")
            canv.restoreState()

        def afterFlowable(self, f):
            lvl = getattr(f, "_toc_level", None)
            if lvl is not None and lvl <= 1:
                self.notify("TOCEntry", (lvl, f.getPlainText(), self.page))

    story = [Spacer(1, 50 * mm),
             Paragraph("FishStocke", ParagraphStyle("t1", parent=H[1], fontSize=30, leading=36, alignment=1)),
             Paragraph(escape(T[lang]["app"]), ParagraphStyle("t2", parent=base, fontSize=12, leading=16, alignment=1)),
             Spacer(1, 15 * mm),
             Paragraph(escape(title), ParagraphStyle("t3", parent=H[1], fontSize=20, leading=26, alignment=1)),
             Spacer(1, 10 * mm),
             Paragraph(f"{T[lang]['date']}: {date.today().strftime('%d.%m.%Y')}",
                       ParagraphStyle("t4", parent=base, alignment=1)),
             PageBreak(),
             Paragraph(T[lang]["toc"], H[1])]
    toc = TableOfContents()
    toc.levelStyles = [ParagraphStyle("toc0", parent=base, fontName="DejaVu-Bold", leftIndent=0),
                       ParagraphStyle("toc1", parent=base, leftIndent=8 * mm),
                       ParagraphStyle("toc2", parent=base, fontSize=9, leftIndent=16 * mm)]
    story.append(toc)
    single = len(parts) == 1 and len(parts[0]["articles"]) == 1
    for part in parts:
        story.append(PageBreak())
        if not single:
            h = Paragraph(escape(part["title"]), H[1])
            h._toc_level = 0
            story.append(h)
        for a in part["articles"]:
            h = Paragraph(escape(a["title"]), H[1 if single else 2])
            h._toc_level = 0 if single else 1
            story.append(h)
            story += flow(_shift(parse(_strip_title(a["markdown"])), 0 if single else 1))
    buf = io.BytesIO()
    try:
        Doc(buf).multiBuild(story)
    finally:
        import shutil
        shutil.rmtree(tmp, ignore_errors=True)
    return buf.getvalue()


# ------------------------------------------------------------------ Markdown одним файлом
def to_markdown(title: str, parts: list[dict]) -> str:
    out = [f"# {title}", "",
           "<!-- Файл создан из справки FishStocke (services/help/app/content). Правьте статьи справки, а не этот файл. -->",
           ""]
    single = len(parts) == 1
    for part in parts:
        if not single:
            out += [f"## {part['title']}", ""]
        for a in part["articles"]:
            body = _strip_title(a["markdown"]).strip("\n")
            body = re.sub(r"\[([^\]]+)\]\(help:[\w-]+\)", r"«\1»", body)   # ссылки справки → название статьи
            shift = 1 if single else 2
            lines = []
            in_code = False
            for line in body.splitlines():
                if line.startswith("```"):
                    in_code = not in_code
                if not in_code and line.startswith("#"):
                    line = "#" * shift + line
                lines.append(line)
            out += ["#" * (shift + 1) + " " + a["title"], "", *lines, ""]
    return "\n".join(out).rstrip() + "\n"
