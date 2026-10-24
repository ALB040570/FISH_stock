import io
import json

from conftest import H, REPO, load_service

CONTENT = REPO / "services" / "help" / "app" / "content"
def _png():
    from PIL import Image
    buf = io.BytesIO()
    Image.new("RGB", (40, 20), (40, 120, 160)).save(buf, "PNG")
    return buf.getvalue()


PNG = _png()


def test_bundle_complete():
    """Каждая статья поставки есть на русском и казахском, первая строка — заголовок."""
    arts = json.loads((CONTENT / "articles.json").read_text(encoding="utf-8"))
    secs = {s["code"] for s in json.loads((CONTENT / "sections.json").read_text(encoding="utf-8"))}
    assert len({a["id"] for a in arts}) == len(arts)
    for a in arts:
        assert a["section"] in secs, a
        for lang in ("ru", "kk"):
            p = CONTENT / lang / f"{a['id']}.md"
            assert p.exists(), p
            assert p.read_text(encoding="utf-8").startswith("# "), p


def test_toc_view_and_rights(svc):
    c = svc("help")
    toc = c.get("/sections?lang=ru", headers=H("viewer")).json()
    assert [s["code"] for s in toc][:8] == ["general", "ichthyologist", "methodology", "mathematician", "dashboards",
                                            "admin", "install", "glossary"]
    kk = c.get("/sections?lang=kk", headers=H("viewer")).json()
    assert kk[0]["title"] == "Жалпы мәліметтер" and all(not a["fallback"] for s in kk for a in s["articles"])
    a = c.get("/articles/data-upload?lang=ru", headers=H("viewer")).json()
    assert a["title"] and a["can_edit"] is False and a["texts"] is None
    assert c.get("/sections").status_code == 401
    body = {"lang": "ru", "markdown": "# Загрузка\n\nНовый текст.", "base_version": a["current_version"]}
    assert c.put("/articles/data-upload", headers=H("ichthyologist"), json=body).status_code == 403
    assert c.put("/articles/data-upload", headers=H("viewer"), json=body).status_code == 403


def test_edit_versions_conflict_restore(svc):
    c = svc("help")
    a = c.get("/articles/map?lang=ru", headers=H("mathematician")).json()
    assert a["can_edit"] and a["texts"]["kk"]
    v0 = a["current_version"]
    r = c.put("/articles/map", headers=H("mathematician"),
              json={"lang": "ru", "markdown": "# Карта\n\nИзменено.", "base_version": v0}).json()
    assert r["changed"] and r["version"] == v0 + 1
    # чужая правка поверх устаревшей версии
    r = c.put("/articles/map", headers=H("admin"), json={"lang": "ru", "markdown": "# Карта\n\nЕщё.", "base_version": v0})
    assert r.status_code == 409
    # без заголовка нельзя
    assert c.put("/articles/map", headers=H("admin"), json={"lang": "ru", "markdown": "текст"}).status_code == 422
    a2 = c.get("/articles/map?lang=ru", headers=H()).json()
    assert a2["edited"] and "Изменено" in a2["markdown"]
    assert a2["texts"]["kk"] == a["texts"]["kk"]          # казахский текст не затронут
    assert c.post(f"/articles/map/restore/{v0}", headers=H("admin")).status_code == 200
    assert c.get("/articles/map?lang=ru", headers=H()).json()["markdown"] == a["markdown"]
    vers = c.get("/articles/map/versions", headers=H()).json()
    assert [v["n"] for v in vers][:3] == [v0 + 2, v0 + 1, v0]
    old = c.get(f"/articles/map?lang=ru&version={v0 + 1}", headers=H()).json()
    assert "Изменено" in old["markdown"]


def test_bundle_update_keeps_user_edits(tmp_path, monkeypatch):
    mod = load_service("help", tmp_path, monkeypatch)
    from fastapi.testclient import TestClient
    c = TestClient(mod.app)
    c.put("/articles/map", headers=H("admin"), json={"lang": "ru", "markdown": "# Карта\n\nМоя правка."})
    orig = mod.bundle_text

    def changed(aid):
        t = orig(aid)
        if aid in ("map", "roles"):
            t["ru"] += "\nНовый абзац поставки.\n"
        return t
    monkeypatch.setattr(mod, "bundle_text", changed)
    mod._seed()
    roles = c.get("/articles/roles?lang=ru", headers=H()).json()
    assert "Новый абзац поставки" in roles["markdown"] and not roles["bundle_update"]     # не правили — обновлена
    m = c.get("/articles/map?lang=ru", headers=H()).json()
    assert "Моя правка" in m["markdown"] and m["bundle_update"]                           # правили — сохранена
    c.post("/articles/map/take-bundled", headers=H("mathematician"))
    m = c.get("/articles/map?lang=ru", headers=H()).json()
    assert "Новый абзац поставки" in m["markdown"] and not m["edited"] and not m["bundle_update"]


def test_search_and_context(svc):
    c = svc("help")
    res = c.get("/search", params={"q": "шаблон Excel", "lang": "ru"}, headers=H("viewer")).json()
    assert res and res[0]["id"] in ("data-upload", "dataset-types", "data-errors")
    assert c.get("/search", params={"q": "үлгі", "lang": "kk"}, headers=H()).json()
    ctx = lambda r: c.get("/context", params={"route": r, "lang": "ru"}, headers=H()).json()["id"]
    assert ctx("/data") == "data-upload"
    assert ctx("/data/01ABC") == "data-edit"
    assert ctx("/calculations/01ABC") == "calculations"   # страница расчёта (решение № 70)
    assert ctx("/catalog") == "catalog" and ctx("/profiles") == "profiles"
    assert ctx("/") == "map"
    assert ctx("/somewhere") == "intro"


def test_new_article_and_image(svc):
    c = svc("help")
    assert c.post("/articles", headers=H("mathematician"),
                  json={"id": "x", "section": "general", "title_ru": "Т"}).status_code == 422
    assert c.post("/articles", headers=H("mathematician"),
                  json={"id": "lab-notes", "section": "mathematician", "order": 99, "title_ru": "Заметки"}).status_code == 201
    img = c.post("/images", headers=H("mathematician"), files={"file": ("shot.png", PNG, "image/png")}).json()
    assert img["markdown"].startswith("![shot](help-image:")
    assert c.post("/images", headers=H("mathematician"), files={"file": ("a.txt", b"x", "text/plain")}).status_code == 415
    assert c.post("/images", headers=H("mathematician"), files={"file": ("b.png", b"not a png", "image/png")}).status_code == 415
    md = f"# Заметки\n\nТекст.\n\n{img['markdown']}\n"
    assert c.put("/articles/lab-notes", headers=H("mathematician"), json={"lang": "ru", "markdown": md}).status_code == 200
    assert c.get(f"/images/{img['name']}", headers=H("viewer")).content == PNG
    r = c.get("/export.docx?lang=ru&article=lab-notes", headers=H("viewer"))
    assert r.status_code == 200
    from docx import Document
    d = Document(io.BytesIO(r.content))
    assert len(d.inline_shapes) == 1
    assert c.delete("/articles/map", headers=H("admin")).status_code == 400        # статья поставки
    assert c.delete("/articles/lab-notes", headers=H("admin")).status_code == 200


def test_export_all(svc):
    from docx import Document
    c = svc("help")
    for lang in ("ru", "kk"):
        r = c.get(f"/export.docx?lang={lang}", headers=H("viewer"))
        assert r.status_code == 200 and "attachment" in r.headers["content-disposition"]
        d = Document(io.BytesIO(r.content))
        heads = [p.text for p in d.paragraphs if p.style.name.startswith("Heading")]
        toc = c.get(f"/sections?lang={lang}", headers=H()).json()
        assert all(s["title"] in heads for s in toc)
        p = c.get(f"/export.pdf?lang={lang}", headers=H("viewer"))
        assert p.status_code == 200 and p.content[:4] == b"%PDF" and len(p.content) > 30000
    s = c.get("/export.pdf?lang=kk&section=install", headers=H()).content
    assert s[:4] == b"%PDF"
    assert c.get("/export.pdf?lang=ru&section=nope", headers=H()).status_code == 404


def test_install_doc_in_sync(svc):
    """docs/install.md — раздел «Установка и сопровождение» справки (tools/gen_install_doc.py)."""
    c = svc("help")
    md = c.get("/export.md?lang=ru&section=install", headers=H()).text
    assert (REPO / "docs" / "install.md").read_text(encoding="utf-8") == md


def test_presenter_overview_only(svc):
    """«Презентация» (№ 62): только обзорные статьи, без формул и установки; ссылки на закрытые статьи — текстом."""
    c = svc("help")
    P = H("presenter")
    toc = c.get("/sections", headers=P).json()
    ids = {a["id"] for s in toc for a in s["articles"]}
    assert ids == {"intro", "interface", "roles", "map", "workflow", "method-scheme"}
    assert all(s["articles"] for s in toc)                                     # пустых разделов нет
    for aid in ("method-overview", "method-bayes", "install-docker", "verification", "stock-trust"):
        assert c.get(f"/articles/{aid}", headers=P).status_code == 404, aid
        assert c.get(f"/articles/{aid}/versions", headers=P).status_code == 404, aid
    sch = c.get("/articles/method-scheme", headers=P, params={"expand": 1}).json()["markdown"]
    assert "$" not in sch and "F = k" not in sch and "bundled:" not in sch       # без формул
    wf = c.get("/articles/workflow", headers=P, params={"expand": 1}).json()["markdown"]
    assert "help:method-overview" not in wf and "help:data-upload" not in wf and "Методика расчёта" in wf
    assert all(x["id"] in ids for x in c.get("/search", headers=P, params={"q": "уловистость"}).json())
    assert c.get("/context", headers=P, params={"route": "/runs/1"}).json()["id"] in ids
    md = c.get("/export.md", headers=P).text
    assert "Байесовская когортная модель" in md and "$" not in md
    # остальные роли видят всё
    full = {a["id"] for s in c.get("/sections", headers=H("viewer")).json() for a in s["articles"]}
    assert full > ids and "method-overview" in full
