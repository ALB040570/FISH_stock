"""
Выгрузка виджетов: Excel (данные на экране) и PDF (рисунок + таблица), отчёт раздела — один PDF.
Содержимое присылает интерфейс (то, что видит пользователь); здесь только оформление.
В каждом файле — водоём, фильтры, источник и версия данных, дата выгрузки.
"""
from __future__ import annotations

import base64
import io
from datetime import datetime
from pathlib import Path

FONT_DIR = Path(__file__).parent / "fonts"
_fonts = False


def _cell(v):
    if isinstance(v, bool):
        return "да" if v else "нет"
    return v


def _src_line(w: dict) -> str:
    parts = []
    for s in w.get("sources") or []:
        parts.append(f"{s.get('title')} (версия {s.get('version')})")
    out = "Источник: " + "; ".join(parts) if parts else ""
    if w.get("filter_label"):
        out += (" · " if out else "") + str(w["filter_label"])
    return out


def to_xlsx(w: dict) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    wb = Workbook()
    ws = wb.active
    ws.title = "Данные"
    cols = w.get("columns") or []
    ws.append([c.get("title") + (f", {c['unit']}" if c.get("unit") and c["unit"] not in str(c.get("title")) else "")
               for c in cols])
    for c in ws[1]:
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="DCE8F2")
        c.alignment = Alignment(wrap_text=True, vertical="top")
    for r in w.get("rows") or []:
        ws.append([_cell(r.get(c["code"])) for c in cols])
    for i, c in enumerate(cols, start=1):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = max(10, min(40, len(str(c.get("title"))) + 4))
    ws.freeze_panes = "A2"
    info = wb.create_sheet("Описание")
    for line in [w.get("title") or "", w.get("subtitle") or "", w.get("context") or "", _src_line(w),
                 *(w.get("notes") or []), f"Выгружено из FishStocke: {datetime.now().strftime('%d.%m.%Y %H:%M')}"]:
        if line:
            info.append([line])
    info.column_dimensions["A"].width = 120
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _styles():
    global _fonts
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    if not _fonts:
        pdfmetrics.registerFont(TTFont("DejaVu", str(FONT_DIR / "DejaVuSans.ttf")))
        pdfmetrics.registerFont(TTFont("DejaVu-Bold", str(FONT_DIR / "DejaVuSans-Bold.ttf")))
        _fonts = True
    ss = getSampleStyleSheet()
    return {
        "h0": ParagraphStyle("h0", parent=ss["Title"], fontName="DejaVu-Bold", fontSize=16, leading=20),
        "h": ParagraphStyle("h", parent=ss["Heading2"], fontName="DejaVu-Bold", fontSize=12.5, leading=16, spaceBefore=10),
        "p": ParagraphStyle("p", parent=ss["Normal"], fontName="DejaVu", fontSize=9, leading=12),
        "s": ParagraphStyle("s", parent=ss["Normal"], fontName="DejaVu", fontSize=7.5, leading=10, textColor="#5B6B74"),
        "t": ParagraphStyle("t", parent=ss["Normal"], fontName="DejaVu", fontSize=7, leading=8.5),
    }


def _fmt(v):
    if v is None:
        return "—"
    if isinstance(v, bool):
        return "да" if v else "нет"
    if isinstance(v, float):
        a = abs(v)
        d = 0 if a >= 100 else 2 if a >= 1 else 3
        return f"{v:,.{d}f}".replace(",", " ").replace(".", ",")
    return str(v)


def _flow(w: dict, S, width, max_rows=60):
    from reportlab.lib import colors
    from reportlab.platypus import Image, Paragraph, Spacer, Table, TableStyle
    from xml.sax.saxutils import escape
    out = [Paragraph(escape(w.get("title") or ""), S["h"])]
    if w.get("subtitle"):
        out.append(Paragraph(escape(w["subtitle"]), S["s"]))
    img = w.get("image")
    if img and img.startswith("data:image/png;base64,"):
        raw = base64.b64decode(img.split(",", 1)[1])
        from reportlab.lib.utils import ImageReader
        ir = ImageReader(io.BytesIO(raw))
        iw, ih = ir.getSize()
        h = width * ih / iw if iw else 200
        out += [Spacer(1, 4), Image(io.BytesIO(raw), width=width, height=min(h, 330))]
    cols = (w.get("columns") or [])[:14]
    rows = (w.get("rows") or [])[:max_rows]
    if cols and rows:
        data = [[Paragraph(escape(str(c.get("title"))), S["t"]) for c in cols]]
        data += [[Paragraph(escape(_fmt(r.get(c["code"]))), S["t"]) for c in cols] for r in rows]
        t = Table(data, repeatRows=1, colWidths=[width / len(cols)] * len(cols))
        t.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.3, colors.HexColor("#C9D4DA")),
                               ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#E6EEF2")),
                               ("VALIGN", (0, 0), (-1, -1), "TOP")]))
        out += [Spacer(1, 4), t]
        if len(w.get("rows") or []) > max_rows:
            out.append(Paragraph(f"Показаны первые {max_rows} строк из {len(w['rows'])}; полностью — в выгрузке Excel", S["s"]))
    for n in w.get("notes") or []:
        out.append(Paragraph(escape(n), S["s"]))
    if _src_line(w):
        out.append(Paragraph(escape(_src_line(w)), S["s"]))
    return out


def to_pdf(title: str, context: str, widgets: list[dict]) -> bytes:
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.units import mm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer
    from xml.sax.saxutils import escape
    S = _styles()
    buf = io.BytesIO()
    page = landscape(A4)
    doc = SimpleDocTemplate(buf, pagesize=page, leftMargin=14 * mm, rightMargin=14 * mm, topMargin=12 * mm,
                            bottomMargin=12 * mm, title=title)
    width = page[0] - 28 * mm

    def footer(canvas, d):
        canvas.saveState()
        canvas.setFont("DejaVu", 7)
        canvas.drawString(14 * mm, 7 * mm, f"FishStocke · {title} · {datetime.now().strftime('%d.%m.%Y %H:%M')}")
        canvas.drawRightString(page[0] - 14 * mm, 7 * mm, str(d.page))
        canvas.restoreState()
    story = [Paragraph(escape(title), S["h0"])]
    if context:
        story.append(Paragraph(escape(context), S["p"]))
    story.append(Spacer(1, 6))
    for w in widgets:
        story += _flow(w, S, width)
        story.append(Spacer(1, 8))
    doc.build(story, onFirstPage=footer, onLaterPages=footer)
    return buf.getvalue()
