"""Представления исходных данных мониторинга: числа сверяются с независимым расчётом по файлам Жайсана."""
import io
import math
import statistics as st

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from conftest import H, REPO, load_service, upload_samples

SAMPLES = REPO / "samples" / "zhaysan"


@pytest.fixture()
def dash(data_stack, tmp_path, monkeypatch):
    dm, da, _ = data_stack
    upload_samples(da)
    mod = load_service("dashboards", tmp_path, monkeypatch)
    for k, cl in (("data_admin", da), ("data_model", dm), ("algo_admin", da)):
        mod.SVC[k] = mod.Svc(k, http=cl)
    return TestClient(mod.app), dm, da


def view(c, name, species=None, **opt):
    r = c.post("/widgets/data", headers=H("viewer"), json={"source": {"kind": "insight", "view": name},
                                                           "filters": {"water_body": "zhaysan", "species": species, **opt.pop("_f", {})},
                                                           "options": opt})
    assert r.status_code == 200, r.text
    return r.json()


def raw_catches():
    ws = load_workbook(SAMPLES / "catches.xlsx", read_only=True)["Данные"]
    rows = list(ws.iter_rows(values_only=True))
    head = [str(h).replace(" *", "") for h in rows[0]]
    return [dict(zip(head, r)) for r in rows[1:]]


def test_cpue_matches_manual(dash):
    c, _, _ = dash
    j = view(c, "cpue", "lesh", gear_kind="net", catch_type="scientific")
    hauls = [r for r in raw_catches() if r["Орудие лова"] == "сеть" and r["Тип лова"] == "научный" and r["Год"] == 2019
             and r["лещ"] is not None]
    vals = [r["лещ"] / (r["Размер орудия (длина), м"] / 100) for r in hauls]
    m = sum(vals) / len(vals)
    se = st.stdev(vals) / math.sqrt(len(vals))
    row = next(r for r in j["rows"] if r["year"] == 2019)
    assert row["n"] == len(vals)
    assert abs(row["mean"] - m) < 1e-3 and abs(row["lo"] - (m - 1.96 * se)) < 1e-3 and abs(row["hi"] - (m + 1.96 * se)) < 1e-3
    assert j["chart"]["series"][0]["lo"] and j["sources"][0]["dataset_type"] == "catches"
    # невод: на гектар облова
    s = view(c, "cpue", "sudak", gear_kind="seine")
    hs = [r for r in raw_catches() if r["Орудие лова"] == "невод" and r["Тип лова"] == "научный" and r["Год"] == 2017
          and r["судак"] is not None]
    m2 = sum(r["судак"] / r["Обловленная площадь, га"] for r in hs) / len(hs)
    assert abs(next(r for r in s["rows"] if r["year"] == 2017)["mean"] - m2) < 1e-3


def test_catch_views(dash):
    c, _, _ = dash
    raw = [r for r in raw_catches() if r["Тип лова"] == "научный"]
    comp = view(c, "composition")
    for r in comp["rows"]:
        assert abs(sum(v for k, v in r.items() if k not in ("year", "_pred") and v) - 100) < 0.05
    cov = view(c, "coverage_places")
    assert sum(cell[2] for cell in cov["chart"]["cells"]) == len(raw)
    months = view(c, "coverage_months")
    assert sum(cell[2] for cell in months["chart"]["cells"]) == len(raw)
    box = view(c, "haul_box", "lesh")
    y = next(r for r in box["rows"] if r["year"] == 2018)
    v = sorted(r["лещ"] for r in raw if r["Год"] == 2018 and r["лещ"] is not None)
    assert y["n"] == len(v) and y["median"] == round(st.median(v), 1) and y["max"] == v[-1]
    cnt = view(c, "hauls_count", catch_type="all")
    assert sum(sum(s["values"]) for s in cnt["chart"]["series"]) == 690
    tc = view(c, "type_compare", "lesh")
    assert {s["key"] for s in tc["chart"]["series"]} == {"scientific", "commercial"}


def test_age_weight_param_quota_views(dash):
    c, _, _ = dash
    h = view(c, "age_heatmap", "lesh")
    for r in h["rows"]:
        tot = sum(v for k, v in r.items() if k not in ("year", "_n") and v is not None)
        assert abs(tot - 100) < 0.5
    assert h["species_shown"] == "lesh"
    assert view(c, "age_heatmap")["species_shown"]                       # вид выбран автоматически
    ms = view(c, "mature_share", "sudak")
    assert all(0 <= r["share"] <= 100 for r in ms["rows"] if r["share"] is not None)
    g = view(c, "growth", "lesh")
    assert g["chart"]["series"][0]["key"] == "avg" and len(g["chart"]["series"]) == 6
    d = view(c, "weight_dev", "lesh")
    assert d["chart"]["fmt"] == "dev"
    b = view(c, "biomass_structure", "lesh")
    for r in b["rows"]:
        assert abs(sum(v for k, v in r.items() if k != "year" and v) - 100) < 0.5
    lc = view(c, "lifecycle")
    assert len(lc["chart"]["species"]) == 8 and all(a <= b for a, b in zip(lc["chart"]["m0"], lc["chart"]["m1"]))
    assert view(c, "species_cards")["cards"]
    pc = view(c, "params_compare", "lesh")
    assert any(r["_current"] for r in pc["rows"])
    cond = view(c, "conditions")
    assert cond["rows"][0]["lake_area_km2"]
    q = view(c, "quota")
    r0 = q["rows"][0]
    assert r0["year"] == 2015 and r0["lesh"] == 3597
    assert abs(r0["_total"] - sum(v for k, v in r0.items() if k not in ("year", "_total") and v)) < 0.01
    qc = view(c, "quota_change")
    assert qc["chart"]["series"][0]["values"]
    assert view(c, "quota_vs_stock", "lesh")["rows"]
    comp = view(c, "completeness")
    assert all(0 <= v <= 100 for r in comp["rows"] for k, v in r.items() if k != "species" and v is not None)
    ov = view(c, "overview")
    assert ov["tiles"][0]["value"] == "2015–2024" and ov["tiles"][1]["value"] == 690


def test_compare_drill_export(dash):
    c, _, _ = dash
    j = view(c, "cpue", "lesh", gear_kind="net", _f={"compare": [2019, 2024]})
    assert j["chart"]["x"] == [2019, 2024] and len(j["rows"]) == 3 and j["rows"][2]["year"] == "2024 − 2019"
    assert abs(j["rows"][2]["mean"] - (j["rows"][1]["mean"] - j["rows"][0]["mean"])) < 1e-3
    h = view(c, "coverage_places", _f={"compare": [2019, 2024]})
    assert h["chart"]["x"] == [2019, 2024] and "_diff" in [col["code"] for col in h["columns"]]
    # строки, из которых сложилась точка
    r = c.post("/widgets/drill", headers=H("viewer"), json={"source": {"kind": "insight", "view": "cpue"},
               "filters": {"water_body": "zhaysan"}, "options": {"catch_type": "scientific"},
               "where": {"year": 2019, "gear_kind": "net"}, "dataset_type": "catches"}).json()
    n = sum(1 for x in raw_catches() if x["Год"] == 2019 and x["Орудие лова"] == "сеть" and x["Тип лова"] == "научный")
    assert r["total"] == n and r["rows"][0]["gear"] == "сеть" and r["dataset_id"]
    # выгрузка
    body = {"title": "Улов на усилие", "subtitle": "проверка", "columns": j["columns"], "rows": j["rows"][:2],
            "sources": j["sources"], "notes": ["n"], "filter_label": "научный лов"}
    x = c.post("/widgets/export.xlsx", headers=H("viewer"), json=body)
    assert x.status_code == 200
    wb = load_workbook(io.BytesIO(x.content))
    assert wb["Данные"].max_row == 3 and "Улов на усилие" in wb["Описание"]["A1"].value
    from PIL import Image
    buf = io.BytesIO(); Image.new("RGB", (300, 150), (14, 94, 111)).save(buf, "PNG")
    import base64
    img = "data:image/png;base64," + base64.b64encode(buf.getvalue()).decode()
    p = c.post("/widgets/export.pdf", headers=H("viewer"), json={**body, "image": img})
    assert p.status_code == 200 and p.content[:4] == b"%PDF"
    rep = c.post("/reports/section.pdf", headers=H("viewer"), json={"title": "Исходные данные — Уловы", "context": "Зайсан",
                                                                   "widgets": [body, {**body, "image": img}]})
    assert rep.status_code == 200 and rep.content[:4] == b"%PDF"


def test_signals_and_rules(dash):
    c, _, _ = dash
    s = view(c, "signals")
    rules = {x["rule"] for x in s["signals"]}
    assert "cpue_low" in rules or "cpue_extreme" in rules
    assert all(x["section"] for x in s["signals"])
    r = c.get("/signal-rules", headers=H("viewer")).json()
    assert len(r["rules"]) == 18 and not r["can_edit"]   # + 3 правила «Запасов» (24.10.2026)
    body = {"rules": [{"code": x["code"], "enabled": x["code"] != "cpue_extreme", "values": x["values"]} for x in r["rules"]]}
    for x in body["rules"]:
        if x["code"] == "cpue_low":
            x["values"]["pct"] = 100
    assert c.put("/signal-rules", headers=H("viewer"), json=body).status_code == 403
    bad = {"rules": [{"code": "cpue_low", "values": {"pct": 1000}}]}
    assert c.put("/signal-rules", headers=H("ichthyologist"), json=bad).status_code == 422
    assert c.put("/signal-rules", headers=H("ichthyologist"), json=body).status_code == 200
    s2 = view(c, "signals")
    assert not {"cpue_low", "cpue_extreme"} & {x["rule"] for x in s2["signals"]}
    got = c.get("/signal-rules", headers=H()).json()
    assert next(x for x in got["rules"] if x["code"] == "cpue_low")["values"]["pct"] == 100
    assert len(c.get("/signal-rules/versions", headers=H()).json()) == 1
    q = view(c, "quality")
    levels = [x["level"] for x in q["rows"]]
    assert levels == sorted(levels, key=lambda v: {"error": 0, "warning": 1, "info": 2}[v])


def test_unknown_view_and_no_data(dash):
    c, _, _ = dash
    r = c.post("/widgets/data", headers=H(), json={"source": {"kind": "insight", "view": "nope"},
                                                   "filters": {"water_body": "zhaysan"}, "options": {}})
    assert r.status_code == 400
    j = c.post("/widgets/data", headers=H(), json={"source": {"kind": "insight", "view": "cpue"},
                                                   "filters": {"water_body": "balkhash"}, "options": {}}).json()
    assert j["empty"] and "Уловы" in j["message"]


def test_kazakh_texts(dash):
    c, _, _ = dash
    s = view(c, "signals", _f={"lang": "kk"})
    assert s["signals"] and not any("улов на усилие" in x["text"] for x in s["signals"])
    q = view(c, "quality", _f={"lang": "kk"})
    assert not any(r["issue"].startswith(("Малая", "Средний вес")) for r in q["rows"])
    j = c.post("/widgets/data", headers=H(), json={"source": {"kind": "insight", "view": "cpue"},
                                                   "filters": {"water_body": "balkhash", "lang": "kk"}, "options": {}}).json()
    assert j["empty"] and j["message"].startswith("Бұл су айдыны")


def test_traffic_journal_explain(dash):
    c, _, _ = dash
    t = view(c, "traffic")
    codes = [r["_code"] for r in t["rows"]]
    assert codes[0] in ("lesh", "plotva") and len(codes) >= 8
    lesh = next(r for r in t["rows"] if r["_code"] == "lesh")
    cp = view(c, "cpue", "lesh", gear_kind="net")
    assert lesh["spark"] == [r["mean"] for r in cp["rows"] if r["mean"] is not None]
    assert set(lesh["_levels"]) >= {"cpue_change"} and lesh["share"] > 0
    one = view(c, "traffic", "sudak")
    assert [r["_code"] for r in one["rows"]] == ["sudak"]
    j = view(c, "journal", "sudak")
    assert len(j["rows"]) == 690 and j["journal"]["years"][0] == 2015 and "м. Тополев" in j["journal"]["places"]
    sp_cols = [col for col in j["columns"] if col.get("species")]
    assert sp_cols[0]["code"] == "sudak" and sp_cols[0]["selected"]
    assert all(r["area"] is None for r in j["rows"] if r["_gear_kind"] == "net")
    assert j["columns"][0]["title"] == "Год" and j["rows"][0]["date"].count(".") == 2
    # пояснения
    e = view(c, "cpue", "lesh", gear_kind="net")["explain"]
    assert len(e["how"]) >= 3 and e["now"][0].startswith("Лещ: в 2024 г.") and e["next"][0]["widget"] == "c_places"
    for name in ("composition", "coverage_places", "haul_box", "age_heatmap", "growth", "weight_dev", "quota", "quality", "traffic"):
        r = view(c, name, "lesh")
        assert r["explain"]["how"] and r["explain"]["now"], name
    o = view(c, "overview")["explain"]["now"]
    assert any("Замечаний к данным" in x for x in o)
    assert view(c, "overview", _f={"lang": "kk"})["explain"]["now"]
    k = view(c, "cpue", "lesh", gear_kind="net", _f={"lang": "kk"})["explain"]
    assert "жылы" in k["now"][0] and "Сызық" in k["how"][0]
    # замечания режима «Все данные по виду» — только этого вида
    s = view(c, "signals", "okun")
    assert s["signals"] and all("Окунь" in x["text"] or "окунь" in x["text"] for x in s["signals"])


def test_comments(dash):
    c, _, _ = dash
    key = {"board": "monitoring", "widget": "c_cpue_net", "water_body": "zhaysan", "species": "lesh"}
    assert c.post("/comments", headers=H("viewer"), json={**key, "text": "x"}).status_code == 403
    assert c.post("/comments", headers=H("ichthyologist"), json={**key, "text": "  "}).status_code == 422
    a = c.post("/comments", headers=H("ichthyologist", uid="u2", login="ira"), json={**key, "text": "Рост 2024 — новые места", "year": 2024})
    assert a.status_code == 201
    c.post("/comments", headers=H("mathematician", uid="u3"), json={**key, "text": "Проверить пробу"})
    c.post("/comments", headers=H("mathematician", uid="u3"), json={**key, "species": "sudak", "text": "другой вид"})
    cnt = c.get("/comments/counts", headers=H("viewer"), params={"board": "monitoring", "water_body": "zhaysan", "species": "lesh"}).json()
    assert cnt == {"c_cpue_net": 2}
    lst = c.get("/comments", headers=H("viewer"), params=key).json()
    assert [x["text"] for x in lst] == ["Проверить пробу", "Рост 2024 — новые места"] and not lst[0]["can_delete"]
    cid = a.json()["id"]
    assert c.delete(f"/comments/{cid}", headers=H("mathematician", uid="u3")).status_code == 403
    assert c.delete(f"/comments/{cid}", headers=H("ichthyologist", uid="u2")).status_code == 200
    assert c.get("/comments/counts", headers=H(), params={"board": "monitoring", "water_body": "zhaysan", "species": "lesh"}).json() == {"c_cpue_net": 1}
